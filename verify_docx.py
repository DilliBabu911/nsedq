from docx import Document
doc = Document(r"d:\Desktop Backup 1\Dilli\Publications\ESWA_NSEDQ\NSEDQ_ESWA_DilliBabuDevineni_Submission.docx")
paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
tables = len(doc.tables)
shapes = len(doc.inline_shapes)
print(f"Paragraphs:  {len(paras)}")
print(f"Tables:      {tables}")
print(f"Figures:     {shapes}")
print(f"Title OK:    {any('NSEDQ' in p and 'Neuro-Symbolic' in p for p in paras)}")
print(f"Real numbers:{any('77.1' in p for p in paras)}")
print(f"Abstract OK: {any('enterprise data pipeline' in p.lower() for p in paras)}")
print(f"Conclusion:  {any('9. Conclusion' in p for p in paras)}")
import os
docx_path = r"d:\Desktop Backup 1\Dilli\Publications\ESWA_NSEDQ\NSEDQ_ESWA_DilliBabuDevineni_Submission.docx"
sz = round(os.path.getsize(docx_path)/1024)
print(f"File size KB:{sz}")
print("\nSection headings found:")
for p in paras:
    if len(p) < 90 and (p[:2].strip().isdigit() or p.startswith("Abstract") or p.startswith("References") or p.startswith("Acknowledgements")):
        print(f"  {p}")
